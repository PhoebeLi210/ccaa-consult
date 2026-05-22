#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
创建档案服务行业Excel模板文件
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_archive_equipment_list():
    """创建设备清单模板"""
    wb = Workbook()
    ws = wb.active
    ws.title = "设备清单"
    
    # 标题
    ws['A1'] = '档案数字化设备清单'
    ws.merge_cells('A1:G1')
    ws['A1'].font = Font(bold=True, size=16)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    
    # 表头
    headers = ['序号', '设备名称', '型号规格', '数量', '购置日期', '状态', '保管人']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='D5E8F0', end_color='D5E8F0', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # 示例数据行
    example_data = [
        ['1', '高速扫描仪', 'XX-XXXX', '2', '2024-01-15', '正常', '张三'],
        ['2', '平板扫描仪', 'XX-XXXX', '1', '2024-02-20', '正常', '李四'],
        ['3', '电脑', 'XX-XXXX', '5', '2024-01-10', '正常', '王五'],
        ['', '', '', '', '', '', ''],
        ['', '', '', '', '', '', ''],
        ['', '', '', '', '', '', ''],
    ]
    
    for row_idx, row_data in enumerate(example_data, 4):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # 设置列宽
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 12
    
    wb.save('archive_equipment_list_template.xlsx')
    print("✓ 创建设备清单模板")

def create_personnel_capability_record():
    """创建人员能力考核记录模板"""
    wb = Workbook()
    ws = wb.active
    ws.title = "人员能力考核"
    
    # 标题
    ws['A1'] = '档案服务人员能力考核记录'
    ws.merge_cells('A1:F1')
    ws['A1'].font = Font(bold=True, size=16)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    
    # 表头
    headers = ['姓名', '岗位', '考核项目', '考核结果', '考核日期', '考核人']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='D5E8F0', end_color='D5E8F0', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # 示例数据行
    example_data = [
        ['张三', '档案整理员', '档案分类能力', '合格', '2024-03-15', '李四'],
        ['李四', '数字化操作员', '扫描操作技能', '合格', '2024-03-15', '王五'],
        ['王五', '质检员', '质量检查能力', '合格', '2024-03-15', '赵六'],
        ['', '', '', '', '', ''],
        ['', '', '', '', '', ''],
        ['', '', '', '', '', ''],
    ]
    
    for row_idx, row_data in enumerate(example_data, 4):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # 设置列宽
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 12
    
    wb.save('personnel_capability_record_template.xlsx')
    print("✓ 创建人员能力考核记录模板")

def create_archive_receiving_register():
    """创建档案接收登记表模板"""
    wb = Workbook()
    ws = wb.active
    ws.title = "档案接收登记"
    
    # 标题
    ws['A1'] = '档案接收登记表'
    ws.merge_cells('A1:H1')
    ws['A1'].font = Font(bold=True, size=16)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    
    # 基本信息
    ws['A3'] = '接收日期：'
    ws['C3'] = '客户名称：'
    ws['E3'] = '项目名称：'
    
    # 表头
    headers = ['序号', '档号/编号', '题名/名称', '年度', '页数', '档案状态', '移交人', '接收人']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='D5E8F0', end_color='D5E8F0', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # 示例数据行
    example_data = [
        ['1', '2024-001', 'XX档案', '2024', '50', '完好', '客户A', '张三'],
        ['2', '2024-002', 'XX档案', '2024', '30', '完好', '客户A', '张三'],
        ['', '', '', '', '', '', '', ''],
        ['', '', '', '', '', '', '', ''],
        ['', '', '', '', '', '', '', ''],
    ]
    
    for row_idx, row_data in enumerate(example_data, 6):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # 设置列宽
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 10
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 12
    
    wb.save('archive_receiving_register_template.xlsx')
    print("✓ 创建档案接收登记表模板")

def create_archive_handover_register():
    """创建档案交接登记表模板"""
    wb = Workbook()
    ws = wb.active
    ws.title = "档案交接登记"
    
    # 标题
    ws['A1'] = '档案交接登记表'
    ws.merge_cells('A1:H1')
    ws['A1'].font = Font(bold=True, size=16)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    
    # 基本信息
    ws['A3'] = '交接日期：'
    ws['C3'] = '客户名称：'
    ws['E3'] = '项目名称：'
    
    # 表头
    headers = ['序号', '档号/编号', '题名/名称', '年度', '页数', '档案状态', '移交人', '接收人']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='D5E8F0', end_color='D5E8F0', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # 示例数据行
    example_data = [
        ['1', '2024-001', 'XX档案', '2024', '50', '完好', '张三', '客户A'],
        ['2', '2024-002', 'XX档案', '2024', '30', '完好', '张三', '客户A'],
        ['', '', '', '', '', '', '', ''],
        ['', '', '', '', '', '', '', ''],
        ['', '', '', '', '', '', '', ''],
    ]
    
    for row_idx, row_data in enumerate(example_data, 6):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # 设置列宽
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 10
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 12
    
    wb.save('archive_handover_register_template.xlsx')
    print("✓ 创建档案交接登记表模板")

if __name__ == "__main__":
    create_archive_equipment_list()
    create_personnel_capability_record()
    create_archive_receiving_register()
    create_archive_handover_register()
    print("\n所有Excel模板创建完成！")
